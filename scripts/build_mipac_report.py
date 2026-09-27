from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT = Path('/Users/macintosh/IdeaProjects/bulan/MagikLayout')
OUT = ROOT / 'output/docx/MagikLayout_MIPAC_TVET_2026_Laporan_Inovasi.docx'
TMP = ROOT / 'tmp/mipac_report'
TMP.mkdir(parents=True, exist_ok=True)

def font(run, size=12, bold=False, italic=False):
    run.font.name = 'Times New Roman'
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    run.font.size = Pt(size); run.bold = bold; run.italic = italic

def para(doc, text='', align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=12, bold=False, italic=False, before=0, after=6, style=None):
    p = doc.add_paragraph(style=style)
    p.alignment = align
    p.paragraph_format.line_spacing = 1
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    if align in (WD_ALIGN_PARAGRAPH.JUSTIFY, WD_ALIGN_PARAGRAPH.LEFT):
        p.paragraph_format.first_line_indent = Inches(0.3)
    if text:
        font(p.add_run(text), size, bold, italic)
    return p

def heading(doc, text, level=1):
    p = para(doc, text.upper(), align=WD_ALIGN_PARAGRAPH.LEFT, bold=True, before=8, after=4, style=f'Heading {level}')
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.keep_with_next = True
    return p

def page_break(doc):
    p = doc.add_paragraph(); p.add_run().add_break(WD_BREAK.PAGE)

def set_cell(cell, text, bold=False, size=10, fill=None):
    cell.text = ''
    if fill:
        shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); cell._tc.get_or_add_tcPr().append(shd)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.LEFT; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1
    font(p.add_run(str(text)), size, bold)

def borders(table):
    tblPr = table._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e = OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:space'),'0'); e.set(qn('w:color'),'D9D9D9'); b.append(e)
    tblPr.append(b)

def table(doc, headers, rows, widths=None, size=9.5):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = True; borders(t)
    for i,h in enumerate(headers): set_cell(t.rows[0].cells[i], h, True, size, 'D9E2F3')
    trPr = t.rows[0]._tr.get_or_add_trPr(); rep=OxmlElement('w:tblHeader'); rep.set(qn('w:val'),'true'); trPr.append(rep)
    for ridx,row in enumerate(rows):
        cells=t.add_row().cells
        for i,v in enumerate(row): set_cell(cells[i], v, False, size, 'F7F9FC' if ridx%2 else None)
    return t

def add_field(p, instr):
    run=p.add_run(); fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), instr); run._r.addnext(fld)

def make_architecture(path):
    im=Image.new('RGB',(1500,720),'white'); d=ImageDraw.Draw(im)
    try: f=ImageFont.truetype('/System/Library/Fonts/Supplemental/Arial.ttf',28); fb=ImageFont.truetype('/System/Library/Fonts/Supplemental/Arial Bold.ttf',30)
    except: f=fb=None
    boxes=[(70,220,350,480,'Pelajar\nPlayground / Challenges\nAR Lab / AI Studio'),(470,220,750,480,'Enjin deterministik\nLayout + grader\nJava evidence'),(870,120,1160,300,'RAG corpus\nEN / BM\nretrieval + citations'),(870,400,1160,580,'Haiku pilihan\nGuard\nfallback selamat'),(1260,220,1430,480,'Bukti\nujian\nlog / output')]
    for x1,y1,x2,y2,txt in boxes:
        d.rounded_rectangle((x1,y1,x2,y2),18,outline='#555555',width=4,fill='#F7F9FC')
        lines=txt.split('\n'); y=y1+35
        for line in lines:
            bb=d.textbbox((0,0),line,font=fb); d.text(((x1+x2-(bb[2]-bb[0]))/2,y),line,fill='#111111',font=fb); y+=48
    arrows=[((350,350),(470,350)),((750,290),(870,210)),((750,410),(870,490)),((1160,210),(1260,300)),((1160,490),(1260,400))]
    for a,b in arrows:
        d.line((a[0],a[1],b[0],b[1]),fill='#E8590C',width=8)
        d.polygon([(b[0],b[1]),(b[0]-18,b[1]-10),(b[0]-18,b[1]+10)],fill='#E8590C')
    d.text((70,55),'Rajah 1. Seni bina MagikLayout',fill='#111111',font=fb)
    im.save(path)

def main():
    arch=TMP/'architecture.png'; make_architecture(arch)
    doc=Document(); sec=doc.sections[0]; sec.page_width=Inches(8.27); sec.page_height=Inches(11.69)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
    styles=doc.styles; styles['Normal'].font.name='Times New Roman'; styles['Normal'].font.size=Pt(12)
    styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman')
    for s in ('Heading 1','Heading 2','Heading 3'):
        styles[s].font.name='Times New Roman'; styles[s].font.size=Pt(12); styles[s].font.bold=True; styles[s].font.color.rgb=RGBColor(0,0,0)
        styles[s]._element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman')
    # Header/footer
    hp=sec.header.paragraphs[0]; hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT; font(hp.add_run('LAPORAN INOVASI MIPAC TVET 2026'),9,True)
    fp=sec.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER; font(fp.add_run('MagikLayout  |  '),9); add_field(fp,'PAGE')
    # Cover
    para(doc,'LAPORAN PERTANDINGAN MIPAC TVET 2026',WD_ALIGN_PARAGRAPH.CENTER,14,True,False,90,18)
    para(doc,'MAGIKLAYOUT',WD_ALIGN_PARAGRAPH.CENTER,24,True,False,20,8)
    para(doc,'Platform pembelajaran interaktif untuk memahami tingkah laku Java Swing melalui enjin deterministik, cabaran berperingkat, Artificial Intelligence (AI) Coach dan Augmented Reality (AR)',WD_ALIGN_PARAGRAPH.CENTER,14,False,False,0,50)
    para(doc,'Kategori Penyertaan: Immersive Digital Teaching Aid',WD_ALIGN_PARAGRAPH.CENTER,12,True,False,0,8)
    para(doc,'Grup Nelang PMU',WD_ALIGN_PARAGRAPH.CENTER,12,False,False,0,4)
    para(doc,'Politeknik Mukah, Sarawak',WD_ALIGN_PARAGRAPH.CENTER,12,False,False,0,4)
    para(doc,'Kursus: DFP50463 Java Based Application Development',WD_ALIGN_PARAGRAPH.CENTER,12,False,False,0,4)
    para(doc,'Disediakan oleh\nMohd Azlan bin Ab Aziz\nHasnah binti Ngah',WD_ALIGN_PARAGRAPH.CENTER,12,False,False,40,4)
    para(doc,'Versi laporan: 14 September 2026',WD_ALIGN_PARAGRAPH.CENTER,11,False,True,60,0)
    page_break(doc)
    heading(doc,'SENARAI KANDUNGAN',1); p=para(doc,'',after=0); add_field(p,'TOC \\o "1-3" \\h \\z \\u'); para(doc,'Nota: klik kanan pada senarai ini dalam Microsoft Word dan pilih Update Field untuk mengemas kini nombor muka surat.',WD_ALIGN_PARAGRAPH.LEFT,9,False,True,8,0)
    page_break(doc)
    heading(doc,'MAKLUMAT KUMPULAN',1)
    group_table = table(doc,['Nama','E-mel','Institusi / Jabatan','Jawatan','Peranan dan sumbangan'],[
        ['Mohd Azlan bin Ab Aziz','[isi e-mel rasmi]','Politeknik Mukah / Jabatan Teknologi Maklumat dan Komunikasi','Pensyarah','Ketua kumpulan; reka bentuk pedagogi, enjin layout, AI governance, pembangunan aplikasi dan penulisan laporan.'],
        ['Hasnah binti Ngah','[isi e-mel rasmi]','Politeknik Mukah / Jabatan Perdagangan','Pensyarah','Penasihat kandungan; semakan bahasa, kesesuaian PdP dan maklum balas pengguna.']],size=9)
    for cell in group_table.rows[1].cells[:2]:
        for run in cell.paragraphs[0].runs: run.italic = True
    para(doc,'Nama dan e-mel ketua kumpulan hendaklah dilengkapkan sebelum penghantaran rasmi. Maklumat dalam kurungan siku ialah medan pentadbiran yang tidak tersedia dalam repositori projek.',WD_ALIGN_PARAGRAPH.LEFT,10,False,True,6,8)
    heading(doc,'ABSTRAK',1)
    abstract=('Sebagai pensyarah yang baru dilantik, saya mendapati pelajar awal diploma sering boleh menulis sintaks Java Swing tetapi masih sukar menjangka apa yang berlaku apabila tetingkap dibesarkan, apabila dua komponen berkongsi satu rantau, atau apabila sebuah JPanel bersarang diperlukan. MagikLayout dibangunkan sebagai ruang latihan interaktif yang menjadikan peraturan layout dapat dilihat, diuji dan dibaiki. Platform ini menggunakan port deterministik algoritma BorderLayout, FlowLayout dan GridLayout; menghasilkan kod Java yang boleh disusun; menyediakan sepuluh cabaran berperingkat; serta menawarkan AI Debugging Studio dengan 12 misi visual. Makmal Augmented Reality menggunakan penjejakan imej untuk tiga misi BorderLayout yang diselaraskan dengan NOSS IT-010-3:2016-C01. AI digunakan secara sempit untuk memberi petunjuk dwibahasa berasaskan korpus yang diluluskan. Enjin menentukan diagnosis dan markah, manakala kod Java, bukti struktur dan fallback kekal deterministik. Pada 14 September 2026, 18 fail ujian dengan 207 ujian, build produksi dan dua mod penilaian RAG sebanyak 138 kes setiap satu lulus. Bukti ini mengesahkan kesiapsiagaan teknikal, bukan keberkesanan pembelajaran yang telah dibuktikan; kajian kelas masih merupakan langkah seterusnya.')
    para(doc,abstract,WD_ALIGN_PARAGRAPH.JUSTIFY,11,False,True,0,4)
    para(doc,'Kata kunci: Java Swing, layout manager, Artificial Intelligence (AI) beretika, pembelajaran interaktif, Augmented Reality (AR), TVET',WD_ALIGN_PARAGRAPH.LEFT,11,False,True,2,8)
    heading(doc,'PENDAHULUAN',1)
    para(doc,'Saya membangunkan MagikLayout daripada masalah yang saya lihat berulang dalam pembelajaran Java GUI: pelajar mengubah saiz atau kedudukan secara cuba-jaya, tetapi tidak benar-benar membina model mental tentang siapa yang menentukan kedudukan komponen. Dalam Swing, komponen tidak diletakkan terus pada piksel oleh pelajar; layout manager mengira ruang, urutan dan hubungan antara komponen. Jika peraturan itu tidak dapat dilihat, kesilapan seperti komponen “hilang” atau berubah tempat terasa seperti tingkah laku rawak.',after=6)
    para(doc,'Tujuan projek ini adalah menyediakan teman rujukan yang boleh dibuka segera dalam pelayar, tanpa akaun dan tanpa pemasangan khas. Skopnya meliputi Playground, Challenges, AI Debugging Studio dan AR Lab. Kumpulan sasaran ialah pelajar tahun pertama dan kedua program diploma yang mengikuti DFP50463 Java Based Application Development, serta pensyarah yang memerlukan demonstrasi dan latihan yang konsisten. Kategori penyertaan ialah Immersive Digital Teaching Aid, dengan elemen AI digunakan sebagai sokongan pembelajaran dan bukan sebagai autoriti penilaian.',after=6)
    heading(doc,'PENYATAAN MASALAH DAN KEBAHARUAN',1)
    para(doc,'Masalah utama ialah jurang antara kod yang ditulis dengan tingkah laku runtime yang dilihat. Penyelesaian biasa seperti slaid, nota teks atau GUI builder menerangkan hasil tetapi sering tidak menunjukkan sebab. MagikLayout mengisi jurang ini dengan memaksa hubungan sebab-akibat: tindakan pengguna mengubah pokok komponen, enjin mengira semula geometri, kanvas memaparkan reflow, dan panel kod menjana Java yang sepadan.',after=6)
    para(doc,'Kebaharuan MagikLayout terletak pada satu enjin yang digunakan oleh semua permukaan. Enjin yang sama memaparkan layout, menilai cabaran, membina bukti Java dan menukar keadaan AR kepada struktur Swing. Ini mengurangkan risiko pelajar melihat satu peraturan dalam visual tetapi menerima peraturan lain ketika penilaian. AI Debugging Studio pula memulakan pembelajaran dengan keadaan rosak dan ramalan pelajar, bukan dengan ruang sembang kosong.',after=6)
    heading(doc,'ANALISIS PENYELESAIAN SEDIA ADA',1)
    table(doc,['Penyelesaian sedia ada','Kekangan','Nilai tambah MagikLayout'],[
        ['Nota / slaid','Peraturan diterangkan secara statik; kesan resize sukar dibayangkan.','Reflow berlaku secara langsung dan boleh diulang.'],
        ['GUI builder','Menyembunyikan sebab dan menghasilkan kebergantungan pada klik.','Menunjukkan pokok komponen, manager dan Java vanilla yang boleh disemak.'],
        ['Chatbot umum','Tidak semestinya melihat struktur sebenar; boleh memberi jawapan tidak tepat atau terlalu awal.','Diagnosis dan markah ditentukan enjin; petunjuk AI dikawal, bersumber dan boleh dinyahaktifkan.'],
        ['Latihan kod biasa','Penilaian sering tertumpu pada output akhir.','Parsons, Reflow, Reverse dan misi visual menilai proses membina model mental.']],size=9)
    heading(doc,'BAHAN DAN KAEDAH INOVASI',1)
    para(doc,'Pembangunan dilaksanakan secara iteratif: mengenal pasti salah faham, memodelkan algoritma JDK, membina permukaan latihan, menambah penilaian deterministik, kemudian menambah sokongan AI dan AR. Teknologi utama ialah React 19, TypeScript, Vite, Three.js dan MindAR. Komponen AR dipecahkan secara code-split supaya permukaan teras tidak dibebankan oleh aset mudah alih.',after=6)
    doc.add_picture(str(arch),width=Inches(6.2)); para(doc,'Rajah 1. Seni bina MagikLayout: enjin deterministik menjadi sumber kebenaran, manakala AI menerangkan dalam ruang yang terhad.',WD_ALIGN_PARAGRAPH.CENTER,10,False,True,2,8)
    para(doc,'Aliran kerja AI ialah: keadaan Swing → gradeReverse dan diagnosis → metadata-first retrieval → pilihan komposisi Claude Haiku → semakan citation, bahasa, kebocoran dan percanggahan → petunjuk atau fallback korpus. Model tidak menentukan lulus/gagal, tidak menjana kod penilaian dan tidak menggantikan Java yang dijana secara deterministik.',after=6)
    table(doc,['Permukaan','Fungsi','Bukti teknikal'],[
        ['Playground','Eksplorasi BorderLayout, FlowLayout, GridLayout, resize, nested panel dan undo/redo.','Enjin layout dan code generator.'],
        ['Challenges','10 latihan dalam mod Parsons, Reflow dan Reverse.','Grader struktur dan geometri deterministik.'],
        ['AI Debugging Studio','12 misi visual: ramal, inspect, repair dan verify.','Korpus dwibahasa, guard, evidence drawer dan fallback.'],
        ['AR Lab','3 misi BorderLayout dengan objek 3D, audio, animasi dan sentuhan.','MindAR, Three.js, Java evidence dan pemetaan NOSS.']],size=9)
    heading(doc,'PENGGUNAAN AI DAN ETIKA AI',1)
    para(doc,'Artificial Intelligence (AI) digunakan untuk mempercepat pembangunan perisian, menyusun draf dokumentasi dan, dalam produk, menyusun satu petunjuk Socratic yang pendek. OpenAI Codex membantu pembangunan dan penyuntingan laporan; Claude Haiku digunakan secara pilihan melalui endpoint yang menerima petikan korpus yang telah dipilih. Peranan manusia kekal pada reka bentuk pedagogi, penulisan korpus, pengesahan teknikal, semakan bahasa dan keputusan tentang kesesuaian kelas.',after=6)
    para(doc,'Korpus ialah set data kecil yang ditulis dan disemak sebagai pasangan Bahasa Inggeris dan Bahasa Malaysia, dengan metadata salah faham, aras petunjuk, sumber, versi dan penyemak. Kejuruteraan prompt menetapkan bahasa keluaran, petikan yang dibenarkan, larangan membocorkan penyelesaian dan keperluan citation; prompt tidak menerima rahsia, nama atau e-mel pelajar. Ketepatan disahkan melalui ujian automatik, semakan terhadap algoritma Swing, pemeriksaan citation dan guard, serta semakan build. Data pelajar tidak menggunakan nama, e-mel atau pengecam peranti; log menggunakan label privasi dan kod sesi. Risiko yang diakui ialah hallucination, bias bahasa, kegagalan rangkaian dan kebocoran jawapan. Sistem menangani risiko ini dengan korpus diluluskan, had aras petunjuk, pemeriksaan bahasa dan percanggahan, fallback teks yang disemak, serta pilihan untuk mematikan AI tanpa menghentikan aktiviti.',after=6)
    heading(doc,'PRAKTIKAL DAN KEBOLEHGUNAAN',1)
    para(doc,'MagikLayout boleh digunakan terus melalui URL awam pada komputer dan telefon. Laluan keyboard, fokus yang jelas, kontras WCAG AA dan reduced-motion support disediakan pada permukaan teras. AR Lab disasarkan untuk Safari iPhone melalui image tracking kerana WebXR tidak tersedia secara seragam dalam persekitaran tersebut. Bahan demonstrasi dan kod QR hendaklah disahkan aktif sebelum penyerahan rasmi.',after=6)
    para(doc,'Dari segi penyelenggaraan, cabaran disimpan sebagai data dan bukan logik yang berselerak; ujian melindungi enjin, grader, korpus, guard dan AR. Dari segi penerimaan pengguna, repositori menyimpan bukti maklum balas formatif, tetapi laporan ini tidak menukar maklum balas itu menjadi tuntutan learning gain. Ujian penerimaan kelas yang dirancang perlu melibatkan masa tugasan, bilangan ralat, skor pra/pasca dan temu bual ringkas.',after=6)
    heading(doc,'IMPAK INOVASI',1)
    table(doc,['Dimensi','Impak yang disokong bukti','Status tuntutan'],[
        ['Keberkesanan / kecekapan','Pelajar boleh melihat sebab reflow dan menerima Java evidence tanpa menukar konteks aplikasi.','Kesan pedagogi masih perlu diuji secara kelas.'],
        ['Kos dan produktiviti','Pelayar mengurangkan keperluan pemasangan; aset dan korpus boleh diguna semula.','Simpanan kos belum dikira dalam RM.'],
        ['Kompetensi TVET','AR Lab memetakan prototaip BorderLayout kepada NOSS IT-010-3:2016-C01.','Sesuai sebagai bukti penyelarasan kurikulum.'],
        ['Keluasan impak','Berpotensi digunakan oleh kursus Java GUI di institusi TVET lain.','Potensi sederhana; belum ada data adopsi luar.']],size=9)
    para(doc,'Bukti teknikal semasa adalah kukuh: 18 fail ujian dan 207 ujian lulus; build produksi lulus; penilaian RAG deterministik dan adversarial masing-masing meliputi 138 kes dengan precision retrieval, top-source accuracy dan coverage 100%, serta contradiction dan leakage 0%. Angka ini membuktikan integriti teknikal komponen yang diuji, bukan kesan pembelajaran atau kadar penerimaan masyarakat.',after=6)
    heading(doc,'PERLUASAN, PENGKOMERSIALAN, KELESTARIAN DAN KEBOLEHSKALAAN',1)
    para(doc,'Perluasan terdekat ialah menjalankan pilot di Politeknik Mukah dengan kohort kecil, mengumpul data pra/pasca dan menambah misi berasaskan miskonsepsi yang benar-benar ditemui. Selepas itu, struktur misi boleh dikembangkan kepada layout manager lain, event handling dan prototaip aplikasi yang berkaitan dengan NOSS. Kerjasama strategik boleh melibatkan pensyarah Java GUI, unit e-pembelajaran dan rakan industri yang menggunakan sistem berasaskan Java.',after=6)
    para(doc,'Model pengkomersialan yang munasabah ialah penyediaan lesson pack tersuai, dashboard pensyarah dan integrasi institusi, tetapi tiada hasil komersial boleh dituntut pada tahap ini. Kelestarian bergantung pada hosting statik, penyelenggaraan korpus, semakan endpoint AI dan ujian regresi. Kebolehskalaan datang daripada seni bina data-driven: cabaran baharu, petikan baharu dan misi baharu boleh ditambah tanpa menulis semula enjin asas.',after=6)
    heading(doc,'KESIMPULAN',1)
    para(doc,'MagikLayout ialah inovasi yang cuba menjawab satu masalah PdP secara praktikal: bagaimana menjadikan peraturan layout Java Swing yang tidak nampak kepada pelajar sebagai sesuatu yang boleh dilihat, diramal, dibaiki dan dibuktikan. Sebagai pensyarah yang baru bermula, saya memilih untuk membina alat yang tidak mengambil alih proses berfikir pelajar. Enjin menentukan kebenaran; AI hanya membantu menerangkan; bukti kod membolehkan semakan. Sistem telah mencapai kesiapsiagaan teknikal untuk demonstrasi, manakala tuntutan learning gain, adoption dan impak skala besar menunggu kajian kelas yang dirancang.',after=6)
    heading(doc,'PENGAKUAN KEASLIAN',1)
    para(doc,'Kami mengesahkan bahawa inovasi dan laporan ini ialah hasil kerja kumpulan kami, tidak memplagiat karya pihak lain dan tidak sengaja melanggar hak cipta, privasi atau kerahsiaan. Semua angka teknikal dalam laporan boleh disemak melalui kod sumber, log ujian dan dokumen sokongan. Medan e-mel rasmi ahli kumpulan perlu dilengkapkan sebelum penghantaran.',after=6)
    heading(doc,'TAHAP KESIAPSEDIAAN TEKNOLOGI',1)
    para(doc,'MagikLayout dicadangkan pada TRL 6: prototaip bersepadu telah dibangunkan, diuji melalui automated suite, build produksi dan penilaian RAG, serta mempunyai laluan AR yang boleh didemonstrasikan. Ia belum diletakkan pada TRL 7-9 kerana validasi bilik darjah, bukti penggunaan merentasi institusi dan proses operasi jangka panjang belum lengkap.',after=6)
    heading(doc,'RUJUKAN',1)
    refs=[
        'Oracle. (2026). Java Platform Standard Edition API documentation: java.awt and javax.swing. https://docs.oracle.com/en/java/javase/21/docs/api/',
        'Qian, Y., & Lehman, J. (2017). Students’ misconceptions and other difficulties in introductory programming: A literature review. ACM Transactions on Computing Education, 18(1), Article 1. https://doi.org/10.1145/3077618',
        'Keuning, H., Jeuring, J., & Heeren, B. (2018). A systematic literature review of automated feedback generation for programming exercises. ACM Transactions on Computing Education, 19(1), Article 3. https://doi.org/10.1145/3231711',
        'CIAST. (2026). Format dan kandungan laporan inovasi pertandingan MIPAC TVET 2026. Pusat Latihan Pengajar dan Kemahiran Lanjutan.',
        'Grup Nelang PMU. (2026). MagikLayout: dokumentasi teknikal, korpus AI, keputusan ujian dan bukti pembangunan [Dokumen dalaman tidak diterbitkan].',
    ]
    for r in refs:
        p=para(doc,r,WD_ALIGN_PARAGRAPH.LEFT,10,False,False,0,4); p.paragraph_format.left_indent=Inches(.25); p.paragraph_format.first_line_indent=Inches(-.25)
    heading(doc,'LAMPIRAN',1)
    para(doc,'Lampiran A: Bukti ujian perisian',WD_ALIGN_PARAGRAPH.LEFT,12,True,False,0,2)
    table(doc,['Item','Keputusan bertarikh 14 September 2026'],[
        ['Test suite','18 test files passed; 207 tests passed.'],
        ['Production build','TypeScript build and Vite production build passed.'],
        ['RAG deterministic gate','138 cases; precision 100%; top-source accuracy 100%; coverage 100%; contradiction 0%; leakage 0%.'],
        ['RAG adversarial gate','138 cases; precision 100%; top-source accuracy 100%; coverage 100%; contradiction 0%; leakage 0%.'],
    ],size=9)
    para(doc,'Lampiran B: Pautan demonstrasi',WD_ALIGN_PARAGRAPH.LEFT,12,True,False,8,2)
    para(doc,'Live build: https://magik-layout.mhdazlan.cc\nPlayground: https://magik-layout.mhdazlan.cc/#/\nChallenges: https://magik-layout.mhdazlan.cc/#/challenges\nAI Debugging Studio: https://magik-layout.mhdazlan.cc/#/classroom\nAR Lab: https://magik-layout.mhdazlan.cc/#/ar-lab',WD_ALIGN_PARAGRAPH.LEFT,10,False,False,0,4)
    para(doc,'Lampiran C: Bahan yang perlu dilengkapkan sebelum submission: e-mel rasmi ahli, surat pengesahan institusi jika diperlukan, gambar penggunaan sebenar, data pilot kelas, pautan video/QR aktif dan bukti kebenaran aset pihak ketiga.',WD_ALIGN_PARAGRAPH.LEFT,10,False,True,8,0)
    doc.core_properties.title='Laporan Inovasi MagikLayout MIPAC TVET 2026'; doc.core_properties.author='Mohd Azlan bin Ab Aziz; Hasnah binti Ngah'
    OUT.parent.mkdir(parents=True, exist_ok=True); doc.save(OUT)
    print(OUT)

if __name__=='__main__': main()
