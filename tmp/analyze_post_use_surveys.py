from openpyxl import load_workbook
from collections import Counter, defaultdict
from statistics import mean, median
import re, json

P1 = "/Users/macintosh/Downloads/kedua LayoutLab Student Evaluation (Responses) (1).xlsx"
P2 = "/Users/macintosh/Desktop/satu MagikLayout Student Experience Survey (Responses).xlsx"


def split_items(value):
    if not value:
        return []
    return [x.strip() for x in re.split(r"(?<=\.),\s+", str(value)) if x.strip()]


ws1 = load_workbook(P1, data_only=True).active
sec1 = defaultdict(list)
sec2 = defaultdict(list)
for r in range(2, ws1.max_row + 1):
    for c, rating in zip(range(2, 7), range(1, 6)):
        for item in split_items(ws1.cell(r, c).value):
            sec1[item].append(rating)
    for c, rating in zip(range(7, 12), range(1, 6)):
        for item in split_items(ws1.cell(r, c).value):
            sec2[item].append(rating)


def short(s):
    repl = {
        "Adding components (clicking or dragging) felt intuitive and predictable.": "Adding components intuitive",
        "Resizing the frame helped me understand how layout managers respond to changing window sizes.": "Resizing aided understanding",
        'The "diff flash" (orange highlight) in the Code Panel made it easy to connect my actions to Java code.': "Diff flash connected actions to code",
        'The two-step confirmation for the "Reset" button prevented accidental data loss.': "Reset confirmation useful",
        "The application was easy to navigate using only my keyboard.": "Keyboard navigation easy",
        "LayoutLab gave me a clearer understanding of Swing managers than trial-and-error in an IDE.": "Clearer than IDE trial-and-error",
        "The generated Java code helped me understand how to write Swing UI code from scratch.": "Generated code aided learning",
        "The Hint Strip and the hidden-component \"vanish\" chips helped me understand Swing\\'s rules.": "Hint strip/vanish chips useful",
        "I trust that the layout behavior I see in LayoutLab matches actual compiled Java code.": "Trusted layout fidelity",
        "Mode 1 (Parsons): Ordering the code magnets helped me understand the correct sequence of Swing initialization.": "Parsons sequencing",
        "Mode 2 (Reflow): Predicting where a component would land after a resize improved my mental model.": "Reflow prediction",
        "Mode 3 (Reverse): Rebuilding a target frame was a fair and effective test of my practical skills.": "Reverse challenge",
        "The grading felt fair because it evaluated structural accuracy/rules, not pixel-perfect errors.": "Structural grading fair",
        "When I made a mistake, the feedback clearly pointed to the rule I needed to fix rather than just giving me a score.": "Feedback identified rule",
    }
    return repl.get(s, s[:70])


def summarize_group(items):
    out = []
    for item, vals in items.items():
        out.append({
            "item": short(item), "n": len(vals), "mean": round(mean(vals), 3),
            "median": median(vals), "top2_n": sum(v >= 4 for v in vals),
            "top2_pct": round(100 * sum(v >= 4 for v in vals) / len(vals), 1),
            "dist": dict(sorted(Counter(vals).items()))
        })
    return sorted(out, key=lambda x: x["mean"], reverse=True)


print("DATASET1_SECTION1")
print(json.dumps(summarize_group(sec1), indent=2))
print("DATASET1_SECTION2")
print(json.dumps(summarize_group(sec2), indent=2))
all1 = [v for vals in list(sec1.values()) + list(sec2.values()) for v in vals]
print("DATASET1_ALL_RATINGS", len(all1), round(mean(all1), 3), round(100*sum(v>=4 for v in all1)/len(all1),1))

ws2 = load_workbook(P2, data_only=True).active
mapping = {"Strongly disagree":1,"Disagree":2,"Neither agree nor disagree":3,"Agree":4,"Strongly agree":5}
labels = [
    "Find activity without help", "Resizing aided understanding", "Code change connected rules to code",
    "Wrong-answer feedback identified reason", "Challenges aided prediction", "AI hint supported repair thinking",
    "Can explain BorderLayout collision", "Can choose layout manager"
]
item2 = []
for c, label in zip(range(4,12), labels):
    raw = [ws2.cell(r,c).value for r in range(2,ws2.max_row+1)]
    vals = [mapping[v] for v in raw if v in mapping]
    item2.append({
        "item":label,"valid_n":len(vals),"not_used":sum(v=="Not used" for v in raw),
        "mean":round(mean(vals),3),"median":median(vals),
        "top2_n":sum(v>=4 for v in vals),"top2_pct":round(100*sum(v>=4 for v in vals)/len(vals),1),
        "dist":dict(sorted(Counter(vals).items()))
    })
print("DATASET2_ITEMS")
print(json.dumps(item2,indent=2))
all2=[]
for c in range(4,12):
    for r in range(2,ws2.max_row+1):
        v=ws2.cell(r,c).value
        if v in mapping: all2.append(mapping[v])
print("DATASET2_ALL_RATINGS",len(all2),round(mean(all2),3),round(100*sum(v>=4 for v in all2)/len(all2),1))

features=Counter()
for r in range(2,ws2.max_row+1):
    value=ws2.cell(r,3).value or ""
    for x in str(value).split(", "):
        if x.strip(): features[x.strip()]+=1
print("FEATURES",json.dumps(features,indent=2))
