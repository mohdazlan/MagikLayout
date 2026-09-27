from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_BREAK
from pathlib import Path


OUT = Path("output/assessment")
OUT.mkdir(parents=True, exist_ok=True)

NAVY = "17365D"
BLUE = "DDEBF7"
PALE = "F3F6FA"
GRAY = "D9D9D9"
MID = "666666"
BLACK = "000000"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color=GRAY, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_width(cell, inches):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, end])


def setup(doc, short_title):
    section = doc.sections[0]
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    section.header_distance = Inches(0.28)
    section.footer_distance = Inches(0.28)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(8)
    title_ppr = title.element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for style_name, size, before, after in [
        ("Heading 1", 16, 12, 6),
        ("Heading 2", 13, 9, 4),
        ("Heading 3", 11, 7, 3),
    ]:
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Small Note" not in styles:
        s = styles.add_style("Small Note", WD_STYLE_TYPE.PARAGRAPH)
        s.font.name = "Aptos"
        s.font.size = Pt(9)
        s.font.color.rgb = RGBColor(70, 70, 70)
        s.paragraph_format.space_after = Pt(4)

    header = section.header.paragraphs[0]
    header.text = short_title
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.name = "Aptos"
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor(90, 90, 90)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("Page ")
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor(90, 90, 90)
    add_page_field(footer)

    doc.core_properties.title = short_title
    doc.core_properties.subject = "Java Swing GUI formative assessment using MagikLayout"
    doc.core_properties.author = "Course lecturer"


def add_title(doc, title, subtitle):
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.add_run(title)
    sub = doc.add_paragraph()
    sub.paragraph_format.space_after = Pt(14)
    run = sub.add_run(subtitle)
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = RGBColor(70, 70, 70)


def add_meta_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, (label, value) in enumerate(rows):
        c0, c1 = table.rows[i].cells
        set_width(c0, 1.55)
        set_width(c1, 5.95)
        set_cell_shading(c0, BLUE)
        for c in (c0, c1):
            set_cell_border(c)
            set_cell_margins(c, 100, 120, 100, 120)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p0 = c0.paragraphs[0]
        p0.add_run(label).bold = True
        c1.paragraphs[0].add_run(value)
    doc.add_paragraph()


def add_bullets(doc, items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)


def add_lines(doc, count=3, label=None):
    if label:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.add_run(label).bold = True
    for _ in range(count):
        p = doc.add_paragraph("________________________________________________________________________________")
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.0
        for r in p.runs:
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(140, 140, 140)


def add_section_heading(doc, text, marks=None):
    heading = text if marks is None else f"{text}  {marks} marks"
    doc.add_heading(heading, level=1)


def add_response_table(doc, headers, widths, rows, font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, h in enumerate(headers):
        c = hdr.cells[i]
        set_width(c, widths[i])
        set_cell_shading(c, NAVY)
        set_cell_border(c)
        set_cell_margins(c, 100, 100, 100, 100)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(9)
    for row_idx, values in enumerate(rows):
        cells = table.add_row().cells
        for i, val in enumerate(values):
            c = cells[i]
            set_width(c, widths[i])
            set_cell_border(c)
            set_cell_margins(c, 110, 110, 110, 110)
            if row_idx % 2 == 1:
                set_cell_shading(c, PALE)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = c.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(font_size)
    doc.add_paragraph()
    return table


def add_checkbox_list(doc, items):
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.12)
        p.paragraph_format.space_after = Pt(3)
        p.add_run("☐  ").font.size = Pt(12)
        p.add_run(item)


def build_student():
    doc = Document()
    setup(doc, "MagikLayout Java GUI Formative Assessment")
    add_title(doc, "Java GUI Programming Formative Assessment", "Using MagikLayout to investigate Java Swing layout managers")
    add_meta_table(doc, [
        ("Student name", "____________________________________________________________"),
        ("Registration number", "____________________________________________________________"),
        ("Class and session", "____________________________   Date: ________________________"),
        ("Assessment", "Individual formative assessment   Duration: 60 minutes   Total: 40 marks"),
    ])

    p = doc.add_paragraph()
    p.add_run("Purpose. ").bold = True
    p.add_run("This assessment checks whether you can predict, build, test, and explain Java Swing layouts. You will record a prediction before using MagikLayout, compare it with the observed result, and explain the rule that caused the result.")

    doc.add_heading("Learning outcomes", level=2)
    add_bullets(doc, [
        "Select BorderLayout, FlowLayout, or GridLayout for a stated interface requirement.",
        "Predict how components move or resize when a container changes size.",
        "Use a nested JPanel to combine layout managers and prevent region conflicts.",
        "Diagnose a layout problem and justify a repair using generated Java evidence.",
    ])

    doc.add_heading("Assessment procedure", level=2)
    add_numbered(doc, [
        "Complete the lecturer's Java Swing Layout Knowledge Pre-Test before opening MagikLayout.",
        "Open https://magik-layout.mhdazlan.cc and complete Tasks 1 to 5 in this booklet.",
        "Write your own prediction before testing each layout. A changed or incorrect prediction is acceptable when your explanation shows what you learned.",
        "Show the requested checkpoints to the lecturer. Submit this booklet before receiving the post-test link.",
        "Complete the post-test once, individually, without returning to MagikLayout or discussing answers.",
    ])

    doc.add_heading("Conditions", level=2)
    add_bullets(doc, [
        "You may use MagikLayout's Playground, Challenges, generated Java panel, and built-in hints when the task permits them.",
        "Do not use search engines, notes, messaging, or another student's screen during the assessment.",
        "The marks are formative: they identify the concepts that need further teaching. Follow your lecturer's instructions about whether the score contributes to coursework.",
    ])

    doc.add_page_break()
    add_section_heading(doc, "Task 1  BorderLayout prediction and observation", 8)
    p = doc.add_paragraph()
    p.add_run("Build. ").bold = True
    p.add_run("In the Playground, select BorderLayout. Add five components labelled Header, Navigation, Workspace, Tools, and Status. Place one component in each available region. Resize the frame wider, narrower, taller, and shorter.")

    add_response_table(doc,
        ["Prompt", "Your prediction before testing", "What you observed in MagikLayout"],
        [2.25, 2.65, 2.65],
        [
            ["Which component will receive most of the extra width and height?", "", ""],
            ["Which components keep a height close to their preferred height?", "", ""],
            ["Which components keep a width close to their preferred width?", "", ""],
        ], font_size=9)

    add_lines(doc, 3, "Explain the layout rule that caused the observed resizing behaviour.")
    add_lines(doc, 2, "Copy one relevant generated Java statement from the Code panel.")
    p = doc.add_paragraph()
    p.add_run("Lecturer checkpoint: ").bold = True
    p.add_run("BorderLayout shown and resized   Initials: __________________   Time: __________")

    doc.add_page_break()
    add_section_heading(doc, "Task 2  FlowLayout and GridLayout comparison", 8)
    p = doc.add_paragraph()
    p.add_run("Build. ").bold = True
    p.add_run("Reset the Playground. Add six buttons labelled A to F. Test the same components first with FlowLayout and then with GridLayout. For each manager, compare a wide frame with a narrow frame.")

    add_response_table(doc,
        ["Feature", "FlowLayout observation", "GridLayout observation"],
        [2.15, 2.7, 2.7],
        [
            ["Component size", "", ""],
            ["Order of components", "", ""],
            ["Effect of narrowing the frame", "", ""],
            ["Use of empty space", "", ""],
        ], font_size=9.2)
    add_lines(doc, 3, "For a row of compact buttons that may wrap, which manager would you choose and why?")
    add_lines(doc, 3, "For a keypad whose buttons must have equal-sized cells, which manager would you choose and why?")

    doc.add_page_break()
    add_section_heading(doc, "Task 3  Repair a BorderLayout region conflict", 10)
    p = doc.add_paragraph()
    p.add_run("Create the problem. ").bold = True
    p.add_run("Use BorderLayout. Add a Save button to SOUTH, then add a Cancel button directly to SOUTH. Record what happens and the rule shown by MagikLayout.")
    add_lines(doc, 3, "What happened to the first button, and why?")

    p = doc.add_paragraph()
    p.add_run("Repair the structure. ").bold = True
    p.add_run("Create a JPanel in SOUTH, give that panel a suitable layout, and place both buttons inside it. Resize the frame to confirm that both remain visible.")

    add_response_table(doc,
        ["Evidence", "Your record"],
        [2.25, 5.3],
        [
            ["Parent container and its layout", ""],
            ["Nested panel and its layout", ""],
            ["Children of the nested panel", ""],
            ["Two generated Java lines that prove the repair", ""],
        ], font_size=9.2)
    add_lines(doc, 3, "Explain why nesting fixes the problem instead of placing both buttons directly in SOUTH.")
    p = doc.add_paragraph()
    p.add_run("Lecturer checkpoint: ").bold = True
    p.add_run("Both buttons visible in the nested panel   Initials: ______________   Time: __________")

    doc.add_page_break()
    add_section_heading(doc, "Task 4  Graded MagikLayout challenges", 8)
    p = doc.add_paragraph()
    p.add_run("Complete two challenges. ").bold = True
    p.add_run("Choose one Reflow challenge and one Reverse challenge. Make one independent attempt before using a hint. Record the evidence below; your lecturer may choose the exact challenge titles.")

    add_response_table(doc,
        ["Evidence", "Challenge 1  Reflow", "Challenge 2  Reverse"],
        [2.0, 2.78, 2.78],
        [
            ["Challenge title", "", ""],
            ["First prediction or structure", "", ""],
            ["First check result", "☐ Correct   ☐ Incorrect", "☐ Correct   ☐ Incorrect"],
            ["Hint used", "☐ No   ☐ Yes", "☐ No   ☐ Yes"],
            ["Final check result", "☐ Passed   ☐ Not yet", "☐ Passed   ☐ Not yet"],
            ["Swing rule learned", "", ""],
        ], font_size=9)
    add_lines(doc, 3, "Choose one mistake you made. Explain how the feedback changed your understanding.")
    p = doc.add_paragraph()
    p.add_run("Lecturer checkpoint: ").bold = True
    p.add_run("Challenge results viewed   Initials: __________________   Time: __________")

    doc.add_page_break()
    add_section_heading(doc, "Task 5  Transfer to a new interface", 6)
    p = doc.add_paragraph()
    p.add_run("Design brief. ").bold = True
    p.add_run("Plan a Student Registration window containing a title, a central area with Name, Programme, and Semester fields, and a bottom row with Submit and Clear buttons. Use only BorderLayout, FlowLayout, GridLayout, and nested JPanel containers.")

    p = doc.add_paragraph()
    p.add_run("A  Draw the component tree or labelled wireframe. ").bold = True
    p.add_run("Show each JPanel, its layout manager, and the component placement.")
    add_lines(doc, 10)

    p = doc.add_paragraph()
    p.add_run("B  Justify your design. ").bold = True
    p.add_run("Explain why each layout manager suits its part of the interface and predict what will happen when the window becomes wider.")
    add_lines(doc, 6)

    doc.add_heading("Submission and post-test", level=1)
    add_checkbox_list(doc, [
        "I wrote my full name and registration number.",
        "I completed all five tasks and obtained the required lecturer checkpoints.",
        "I submitted this hard-copy booklet before opening the post-test.",
        "I completed the post-test once and submitted it successfully.",
    ])
    p = doc.add_paragraph()
    p.add_run("Student declaration: ").bold = True
    p.add_run("The responses in this booklet and the post-test are my own work.")
    doc.add_paragraph("Signature: _________________________________________    Date: __________________")

    p = doc.add_paragraph(style="Small Note")
    p.add_run("Post-use experience survey. ").bold = True
    p.add_run("This questionnaire is separate from the learning assessment. Complete it only according to the consent and research instructions provided by the lecturer; survey responses should not affect course marks.")

    path = OUT / "Java_GUI_MagikLayout_Formative_Assessment_Student.docx"
    doc.save(path)
    return path


def build_guide():
    doc = Document()
    setup(doc, "MagikLayout Formative Assessment Marking Guide")
    add_title(doc, "Java GUI Programming Formative Assessment Marking Guide", "Administration, rubric, feedback, and pre-test post-test alignment")
    add_meta_table(doc, [
        ("Assessment", "Java GUI Programming using MagikLayout"),
        ("Target group", "First- or second-year students learning Java Swing layout managers"),
        ("Recommended duration", "Pre-test 10 minutes; activity 60 minutes; post-test 10 minutes"),
        ("Student evidence", "Completed hard-copy booklet, lecturer checkpoints, and pre-test/post-test submissions"),
    ])

    p = doc.add_paragraph()
    p.add_run("Recommended decision. ").bold = True
    p.add_run("Use the booklet as a low-stakes formative assessment and use the online pre-test and post-test as the learning measure. The assessment deliberately records predictions before observation, requires students to manipulate the layout, and asks them to explain the generated Java. This creates evidence of reasoning rather than attendance alone.")

    doc.add_heading("Learning outcome alignment", level=1)
    add_response_table(doc,
        ["Outcome", "Student evidence", "Task", "Marks"],
        [2.25, 3.55, 0.8, 0.8],
        [
            ["Predict placement and resizing", "Prediction and observation records for BorderLayout", "1", "8"],
            ["Compare layout-manager behaviour", "FlowLayout and GridLayout comparison", "2", "8"],
            ["Build and explain a nested layout", "Conflict repair, component structure, generated Java", "3", "10"],
            ["Respond to deterministic feedback", "Reflow and Reverse challenge records", "4", "8"],
            ["Transfer knowledge to a new design", "Component tree or wireframe and justification", "5", "6"],
        ], font_size=9)

    doc.add_heading("Administration protocol", level=1)
    add_numbered(doc, [
        "Prepare one numbered student booklet for each student. Record the booklet number beside the student's name or registration number.",
        "Release the pre-test first. Require one submission, hide correct answers, and close or remove access when the activity begins.",
        "Give students the booklet and the MagikLayout URL. Read the assessment conditions aloud and start the 60-minute session.",
        "Verify Task 1, Task 3, and Task 4 at the student's screen. Initial only after seeing the required state. The checkpoint is evidence of tool use, not an extra mark.",
        "Collect the booklet before releasing the post-test. Keep students seated and ask them to complete the post-test once without reopening MagikLayout.",
        "Export both Google Form response sheets and match pre-test/post-test records using registration number. Store the assessment marks separately from research-consent data.",
        "Give corrective feedback at the next class. Re-teach concepts with weak item performance or common rubric errors.",
    ])

    doc.add_heading("Recommended controls for the master's study", level=2)
    add_bullets(doc, [
        "Use the same instructions, time limits, device conditions, and lecturer script for every class included in the study.",
        "Use a parallel post-test with the same constructs and difficulty as the pre-test, but change the component names and scenarios. Keep the scoring scheme identical.",
        "Do not show correct answers or item-level feedback between the pre-test and post-test.",
        "Record absences, incomplete attempts, technical failures, and whether a student used a hint. Decide exclusion rules before analysing results.",
        "Analyse paired scores only for students with a valid pre-test and post-test identifier. Report sample size, missing data, score change, and an appropriate paired statistical test with an effect size.",
        "Keep the post-use experience survey separate from course marks unless the institution's ethics approval and consent process explicitly permit a required research questionnaire.",
    ])

    doc.add_page_break()
    doc.add_heading("Analytic marking rubric", level=1)
    p = doc.add_paragraph()
    p.add_run("Scoring principle. ").bold = True
    p.add_run("Award marks for accurate reasoning and evidence. A wrong initial prediction can still earn observation and explanation marks when the student clearly identifies the corrected rule.")

    add_response_table(doc,
        ["Criterion", "Full-credit evidence", "Partial-credit guidance", "Marks"],
        [1.55, 3.0, 2.45, 0.55],
        [
            ["Task 1 prediction", "Three predictions recorded before testing and logically connected to layout behaviour.", "1 mark per meaningful prediction; no credit for blank or copied-after-testing entries.", "3"],
            ["Task 1 observation and rule", "Observations distinguish CENTER, horizontal edge regions, and vertical edge regions; explanation links size changes to BorderLayout.", "Award 1 mark per accurate observation and up to 2 for the explanation/code evidence.", "5"],
            ["Task 2 observations", "Comparison accurately distinguishes preferred-size wrapping from equal-size cells and records narrowing behaviour.", "Award across the four comparison rows according to accuracy and completeness.", "4"],
            ["Task 2 selection", "Chooses a suitable manager for compact wrapping controls and for an equal-cell keypad, with reasons.", "2 marks for each justified choice; 1 for an unsupported correct choice.", "4"],
            ["Task 3 diagnosis", "States that two direct occupants conflict in one BorderLayout region and records the visible result.", "Award for accurate observation even if terminology is incomplete.", "3"],
            ["Task 3 repair", "Nested JPanel is placed in SOUTH and contains both buttons under a suitable child layout.", "Award for a structurally valid repair; deduct when both buttons are still direct frame children.", "4"],
            ["Task 3 explanation", "Component hierarchy and generated Java show why the nested panel preserves both buttons.", "Award for hierarchy, code evidence, and causal explanation.", "3"],
        ], font_size=8.6)

    doc.add_page_break()
    doc.add_heading("Analytic marking rubric continued", level=1)
    add_response_table(doc,
        ["Criterion", "Full-credit evidence", "Partial-credit guidance", "Marks"],
        [1.55, 3.0, 2.45, 0.55],
        [
            ["Task 4 challenge evidence", "Two requested modes attempted, first results and final results recorded, checkpoints verified.", "Award separately for each completed challenge record.", "4"],
            ["Task 4 learning explanation", "Identifies a specific misconception and explains how feedback changed the student's reasoning.", "General praise of the tool without a Swing rule earns limited credit.", "4"],
            ["Task 5 structure", "Wireframe or component tree uses valid nesting and labels layouts and placements.", "Award for a workable partial structure even when one layout choice is weak.", "3"],
            ["Task 5 justification", "Explains layout choices and gives a plausible resize prediction.", "Award 1 mark per defensible explanation element.", "3"],
        ], font_size=8.6)

    doc.add_heading("Suggested performance bands", level=2)
    add_response_table(doc,
        ["Score", "Interpretation", "Recommended feedback"],
        [1.1, 2.6, 3.85],
        [
            ["34–40", "Secure", "Proceed to a more complex nested GUI task; ask the student to explain trade-offs."],
            ["26–33", "Developing", "Review one weak layout manager and repeat one challenge with explanation."],
            ["18–25", "Needs support", "Use a guided demonstration followed by prediction and retesting."],
            ["0–17", "Foundational gaps", "Re-teach container, child, layout manager, region, and resize concepts before reassessment."],
        ], font_size=9)

    doc.add_page_break()
    doc.add_heading("Post-test blueprint", level=1)
    p = doc.add_paragraph()
    p.add_run("Build a parallel 10-item post-test. ").bold = True
    p.add_run("Use new interface scenarios while preserving the pre-test's constructs, number of options, scoring, and approximate reading demand. Randomise option order only when it does not change the meaning.")
    add_response_table(doc,
        ["Items", "Construct", "Observable response", "Cognitive demand"],
        [0.75, 2.25, 3.5, 1.05],
        [
            ["1–2", "BorderLayout regions and resizing", "Predict placement or identify the region that absorbs remaining space.", "Apply"],
            ["3–4", "FlowLayout", "Predict order, wrapping, preferred-size behaviour, or alignment.", "Apply"],
            ["5–6", "GridLayout", "Predict equal-cell sizing, rows, columns, order, or resize behaviour.", "Apply"],
            ["7–8", "Nested panels and region conflict", "Diagnose a missing component or choose a valid nested repair.", "Analyse"],
            ["9–10", "Generated Java and layout selection", "Interpret a short code fragment or choose a manager for a stated requirement.", "Analyse"],
        ], font_size=9)

    doc.add_heading("Data sheet structure", level=2)
    add_response_table(doc,
        ["Field", "Example or coding rule"],
        [2.25, 5.3],
        [
            ["participant_id", "Use the same registration number or a study code in both tests."],
            ["class_batch", "Pilot class, second batch, or another predefined group label."],
            ["pre_total and post_total", "0–10; preserve item-level scores as separate columns."],
            ["assessment_total", "0–40 from the hard-copy rubric."],
            ["hint_use", "Record no/yes or the count if MagikLayout makes it available."],
            ["completion_status", "Complete, incomplete, absent, or technical failure."],
        ], font_size=9.2)

    doc.add_page_break()
    doc.add_heading("Feedback codes for faster marking", level=2)
    add_response_table(doc,
        ["Code", "Meaning", "Feedback to student"],
        [0.8, 2.55, 4.2],
        [
            ["BL1", "Region behaviour unclear", "Resize the frame again and compare CENTER with NORTH or SOUTH."],
            ["BL2", "Two direct occupants in one region", "Place a JPanel in the region, then add both controls to that panel."],
            ["FL1", "Flow wrapping misunderstood", "Record component preferred sizes, then narrow the container gradually."],
            ["GL1", "Grid sizing misunderstood", "Compare every cell after resizing; identify what remains equal."],
            ["NP1", "Hierarchy unclear", "Draw the frame, child panels, and controls as a tree before reading the Java."],
        ], font_size=9)

    doc.add_heading("Rationale and sources", level=1)
    p = doc.add_paragraph()
    p.add_run("Prediction before observation. ").bold = True
    p.add_run("Brown and Wilson recommend asking programming learners to make and record predictions before demonstrations. Tasks 1 and 2 therefore capture a prediction before the student resizes or switches layouts.")

    p = doc.add_paragraph()
    p.add_run("Immediate but interpretable feedback. ").bold = True
    p.add_run("Reviews of automated programming assessment distinguish formative feedback from grades alone and note the value of prompt feedback and repeated attempts. Tasks 3 and 4 require students to interpret MagikLayout's deterministic result and state the rule learned.")

    p = doc.add_paragraph()
    p.add_run("Authentic Swing behaviour. ").bold = True
    p.add_run("The Oracle Swing tutorials describe layout managers as controlling component size and position, identify FlowLayout wrapping, GridLayout equal-sized cells, and BorderLayout's region-based behaviour. The assessment uses these observable behaviours as its content domain.")

    doc.add_heading("References", level=2)
    refs = [
        "Brown, N. C. C., & Wilson, G. (2018). Ten quick tips for teaching programming. PLOS Computational Biology, 14(4), e1006023. https://doi.org/10.1371/journal.pcbi.1006023",
        "Keuning, H., Jeuring, J. T., & Heeren, B. (2018). A systematic literature review of automated feedback generation for programming exercises. ACM Transactions on Computing Education, 19(1), Article 3. https://doi.org/10.1145/3231711",
        "Messer, M., Brown, N. C. C., Kölling, M., & Shi, M. (2024). Automated grading and feedback tools for programming education: A systematic review. ACM Transactions on Computing Education, 24(1), Article 10. https://doi.org/10.1145/3636515",
        "Oracle. (n.d.). Laying out components within a container. The Java Tutorials. https://docs.oracle.com/javase/tutorial/uiswing/layout/",
        "Van Petegem, C., Maertens, R., Strijbol, N., et al. (2023). Dodona: Learn to code with a virtual co-teacher that supports active learning. SoftwareX, 24, 101578. https://doi.org/10.1016/j.softx.2023.101578",
    ]
    for ref in refs:
        p = doc.add_paragraph(ref)
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.first_line_indent = Inches(-0.3)
        p.paragraph_format.space_after = Pt(6)

    doc.add_page_break()
    doc.add_heading("Lecturer planning record", level=1)
    add_lines(doc, 2, "Class and date")
    add_lines(doc, 2, "Pre-test link or QR code")
    add_lines(doc, 2, "Post-test link or QR code")
    add_lines(doc, 2, "Experience survey and consent instructions")
    add_lines(doc, 3, "Changes made after the pilot and reason")

    path = OUT / "Java_GUI_MagikLayout_Formative_Assessment_Marking_Guide.docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    print(build_student())
    print(build_guide())
