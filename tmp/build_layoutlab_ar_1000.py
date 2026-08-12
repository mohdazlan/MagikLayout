from pathlib import Path
import re

from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from build_layoutlab_ar_doc import (
    ROOT,
    TMP_DIR,
    NAVY,
    BLUE,
    TEAL,
    ORANGE,
    GREEN,
    INK,
    MUTED,
    PALE,
    PALE_BLUE,
    PALE_GREEN,
    WHITE,
    add_body,
    add_callout,
    add_caption,
    add_heading,
    add_run,
    add_table,
    make_placeholder,
    rgb,
    set_cell_border,
    set_cell_margins,
    set_cell_shading,
    set_table_widths,
    setup_document,
)


OUT_DIR = ROOT / "output" / "docx"
OUT_FILE = OUT_DIR / "LayoutLab_AR_Extended_Abstract_1000_Words.docx"


def add_figure_pair(doc, items):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_widths(table, [4560, 4560])
    for index, (image_path, caption) in enumerate(items):
        cell = table.cell(0, index)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        set_cell_margins(cell, top=30, start=80, bottom=30, end=80)
        set_cell_border(
            cell,
            top={"val": "nil"},
            bottom={"val": "nil"},
            left={"val": "nil"},
            right={"val": "nil"},
        )
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(3)
        p.add_run().add_picture(str(image_path), width=Inches(2.78))
        cp = cell.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_after = Pt(0)
        add_run(cp, caption, italic=True, color=MUTED, size=8.1)
    return table


def word_count(doc):
    words = []
    for node in doc.element.body.iter():
        if node.tag.endswith("}t") and node.text:
            words.extend(re.findall(r"[A-Za-z0-9]+(?:[-’'][A-Za-z0-9]+)*", node.text))
    return len(words)


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    TMP_DIR.mkdir(parents=True, exist_ok=True)

    specs = [
        ("short-figure-1.png", "SCREENSHOT 1", "Target found + anchored 3D model"),
        ("short-figure-2.png", "SCREENSHOT 2", "CENTER resize test — 2/3"),
        ("short-figure-3.png", "SCREENSHOT 3", "SOUTH X-ray + JPanel repair"),
        ("short-figure-4.png", "SCREENSHOT 4", "3/3 + generated Java evidence"),
    ]
    images = []
    for filename, label, subtitle in specs:
        path = TMP_DIR / filename
        make_placeholder(path, label, subtitle, aspect=(620, 760))
        images.append(path)

    doc = setup_document()
    section = doc.sections[0]
    section.top_margin = Inches(0.66)
    section.bottom_margin = Inches(0.64)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(5)
    add_run(p, "MIPAC TVET 2026 • CIAST", bold=True, color=ORANGE, size=9)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(3)
    add_run(p, "LAYOUTLAB AR", bold=True, color=NAVY, size=23, font="Aptos Display")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(5)
    add_run(
        p,
        "A NOSS-Aligned Augmented Reality Trainer for Java Swing Interface Prototyping",
        bold=True,
        color=BLUE,
        size=14,
        font="Aptos Display",
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(3)
    add_run(p, "Mohd Azlan bin Ab Aziz¹  |  Hasnah binti Ngah²", bold=True, size=10)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    add_run(p, "Politeknik Mukah, Sarawak, Malaysia", color=MUTED, size=8.7)

    add_heading(doc, "Extended Abstract", 1)
    add_body(
        doc,
        "Java Swing layout managers are difficult for novice learners because their rules remain invisible until code is compiled. LayoutLab AR is a browser-based augmented reality teaching aid aligned to IT-010-3:2016-C01, Application Prototype Development. Using an iPhone camera and an original tracked target, learners manipulate a virtual JFrame and complete three assessed BorderLayout missions. They place an application title in NORTH, test that CENTER absorbs remaining space during resizing, and diagnose why two buttons cannot independently occupy SOUTH before repairing the structure with a nested JPanel. Progress depends on touching the tracked 3D objects; conventional answer buttons cannot complete the assessment. The final AR state produces deterministic Java code showing the implemented prototype. The module integrates text, tracked 3D graphics, animation, bilingual audio, spatial interaction and generated code, exceeding the organiser’s requirement for three media elements. Two formative evaluations of the wider LayoutLab browser platform (n=11 and n=6) established the learning need and informed the design; they are not claimed as proof of AR effectiveness. The competition build has completed production acceptance on an iPhone 13, including camera permission, stable tracking, all missions, audio and Java evidence. Seventy-five automated tests protect the engine, mission logic and code-generation path. LayoutLab AR therefore makes AR essential to constructing, testing and repairing an application prototype rather than using it as a decorative visual overlay.",
    )
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    add_run(p, "Keywords: ", bold=True, color=BLUE, size=9)
    add_run(p, "augmented reality, BorderLayout, Java Swing, NOSS, prototype development, TVET", color=MUTED, size=9)

    add_table(
        doc,
        ["Competition alignment", "Application"],
        [
            ("NOSS", "IT-010-3:2016 — Pembangunan Aplikasi, Tahap 3"),
            ("Competency Unit", "IT-010-3:2016-C01 — Application Prototype Development"),
            ("Work Activities", "Implement Application Prototype Mock-Up Flow; Conduct UI/UX Functionality Test"),
        ],
        [2100, 7020],
        font_size=8.8,
    )

    doc.add_page_break()

    add_heading(doc, "1. Problem and Objective", 1)
    add_body(
        doc,
        "In BorderLayout, each region accepts one direct component and CENTER absorbs the remaining space. Learners commonly misread these rules, causing disappearing components and unpredictable resizing. The objective is to let learners construct, functionally test, diagnose and repair a Swing prototype through direct AR manipulation while producing auditable NOSS evidence.",
    )

    add_heading(doc, "2. AR Innovation and Assessed Flow", 1)
    add_body(
        doc,
        "The camera recognises an original LayoutLab target and anchors a colour-coded 3D JFrame to it. Touch coordinates are raycast against mission-relevant meshes, making tracking, spatial anchoring and 3D interaction necessary for completion. BorderLayout is intentionally the assessed focus; FlowLayout and GridLayout remain enrichment content in the wider Playground.",
    )
    add_table(
        doc,
        ["Mission", "AR task", "Evidence"],
        [
            ("1 — Construct", "Touch NORTH to place the application title.", "Correct constraint and updated Java state."),
            ("2 — Test", "Touch CENTER and observe the resize animation.", "Verified dynamic layout behaviour."),
            ("3 — Diagnose and repair", "Reveal the SOUTH collision and activate the JPanel repair.", "Both buttons visible; score 3/3; nested-panel Java."),
        ],
        [1650, 3800, 3670],
        font_size=8.7,
    )
    add_callout(
        doc,
        "Why the AR is essential",
        "The tracked object is the learning model, touch surface and assessment interface. The learner must act on the 3D regions to progress.",
        fill=PALE_GREEN,
        title_color=GREEN,
    )
    add_figure_pair(
        doc,
        [
            (images[0], "Figure 1. Target detected and 3D model anchored."),
            (images[1], "Figure 2. CENTER functionality test."),
        ],
    )

    doc.add_page_break()

    add_figure_pair(
        doc,
        [
            (images[2], "Figure 3. Collision diagnosis and repair."),
            (images[3], "Figure 4. Completion and Java evidence."),
        ],
    )

    add_heading(doc, "3. Media Elements", 1)
    add_body(
        doc,
        "Six elements have defined learning functions: text communicates missions and rules; tracked 3D graphics expose the layout model; animation demonstrates resize and collision behaviour; BM/English audio provides optional guidance; direct touch assesses spatial decisions; and deterministic Java connects the AR prototype to implementation evidence.",
    )

    add_heading(doc, "4. Validation and Competition Readiness", 1)
    add_table(
        doc,
        ["Evidence", "Result", "Interpretation"],
        [
            ("Formative platform evaluation", "n=11 and n=6; visualisation, placement and resizing were valued.", "Supports the need and design—not causal AR effectiveness."),
            ("Technical quality", "75 automated tests and strict production build pass.", "Supports deterministic, reproducible output."),
            ("iPhone 13 acceptance", "HTTPS camera, tracking, direct touch, three missions, audio and Java verified.", "Ready for live competition demonstration."),
        ],
        [2150, 3920, 3050],
        font_size=8.5,
    )
    add_body(
        doc,
        "Before judging, the team should complete a small AR pilot measuring target-acquisition time, mission completion, wrong-region attempts and successful repair. This evidence should be reported separately from the earlier browser surveys.",
    )

    add_heading(doc, "5. Impact and Conclusion", 1)
    add_body(
        doc,
        "LayoutLab AR shortens the distance between an abstract layout rule and its visible consequence. Learners construct a state, test behaviour, diagnose a failure, repair the structure and receive verifiable Java evidence. Lecturers gain a repeatable demonstrator aligned to a defined competency, while web deployment avoids native installation. The innovation is strongest when presented as a focused BorderLayout competency module: AR is the method of learning and assessment, and deterministic code—not generative AI—is the technical ground truth.",
    )

    add_heading(doc, "Declaration and Reference", 1)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    add_run(
        p,
        "The work is original and its evidence can be verified. Generative AI does not determine correctness, scoring or assessed Java output.",
        color=INK,
        size=8.4,
    )
    refs = [
        "Keuning, H., Jeuring, J., & Heeren, B. (2018). Automated feedback generation for programming exercises. ACM TOCE, 19(1).",
    ]
    for ref in refs:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.first_line_indent = Inches(-0.18)
        p.paragraph_format.space_after = Pt(2)
        add_run(p, ref, color=MUTED, size=7.8)

    props = doc.core_properties
    props.title = "LayoutLab AR — Extended Abstract (Maximum 1000 Words)"
    props.subject = "MIPAC TVET 2026 CIAST"
    props.author = "Mohd Azlan bin Ab Aziz; Hasnah binti Ngah"

    count = word_count(doc)
    if count > 1000:
        raise RuntimeError(f"Document exceeds limit: {count} words")
    doc.save(OUT_FILE)
    print(f"{OUT_FILE}\nWord count: {count}")


if __name__ == "__main__":
    main()
